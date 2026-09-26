import os
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter
from tokenizers import Tokenizer, models, pre_tokenizers

from app.modules.ingestion.chunking import load_tokenizer, split_sections
from app.modules.ingestion.parsers import ParseError, Section, parse_file


@pytest.mark.parametrize("suffix", ["txt", "md"])
def test_text_formats_preserve_lines(tmp_path, suffix):
    path = tmp_path / ("file." + suffix)
    path.write_text("校园图书馆\n周一开放。", encoding="utf-8")
    sections = parse_file(path, suffix)
    assert [s.text for s in sections] == ["校园图书馆", "周一开放。"]
    assert sections[1].line_number == 2


def test_docx_chinese_and_paragraph(tmp_path):
    document = Document()
    document.add_paragraph("校园图书馆")
    document.add_paragraph("周一开放。")
    path = tmp_path / "file.docx"
    document.save(path)
    assert parse_file(path, "docx")[1].paragraph_number == 2


def test_pdf_text_and_page(tmp_path):
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(200, 200)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 10 10 Td (Library opens Monday) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    path = tmp_path / "file.pdf"
    writer.write(path)
    sections = parse_file(path, "pdf")
    assert "Library" in sections[0].text
    assert sections[0].page_number == 1


@pytest.mark.parametrize(
    "suffix,body,code", [("txt", b"", "NO_TEXT"), ("docx", b"broken", "INVALID_FORMAT")]
)
def test_invalid_files_have_safe_errors(tmp_path, suffix, body, code):
    path = tmp_path / ("file." + suffix)
    path.write_bytes(body)
    with pytest.raises(ParseError) as error:
        parse_file(path, suffix)
    assert error.value.code == code


def test_scanned_pdf_rejected(tmp_path):
    writer = PdfWriter()
    writer.add_blank_page(200, 200)
    path = tmp_path / "scan.pdf"
    writer.write(path)
    with pytest.raises(ParseError, match="NO_TEXT"):
        parse_file(path, "pdf")


def test_word_level_tokenizer_bound_and_overlap():
    tokenizer = Tokenizer(
        models.WordLevel({"[UNK]": 0, **{str(i): i + 1 for i in range(900)}}, unk_token="[UNK]")
    )
    tokenizer.pre_tokenizer = pre_tokenizers.Whitespace()
    pieces = split_sections(
        [Section(" ".join(str(i) for i in range(900)), line_number=1)], tokenizer
    )
    encoded = [tokenizer.encode(p.text, add_special_tokens=False).ids for p in pieces]
    assert list(map(len, encoded)) == [400, 400, 200]
    assert encoded[0][-50:] == encoded[1][:50]
    assert pieces[0].line_number == 1


@pytest.fixture
def real_bge_tokenizer():
    path = os.environ.get("EMBEDDING_TOKENIZER_PATH")
    if not path:
        pytest.skip("EMBEDDING_TOKENIZER_PATH is required for real BGE integration")
    assert Path(path).is_file(), "Configured real BGE tokenizer does not exist"
    return load_tokenizer(path)


def test_real_bge_chunks_stay_bounded_after_independent_encoding(real_bge_tokenizer):
    text = "中 " * 345 + "unaffordability " + "中 " * 500
    pieces = split_sections([Section(text, page_number=3)], real_bge_tokenizer)
    lengths = [
        len(real_bge_tokenizer.encode(piece.text, add_special_tokens=False).ids) for piece in pieces
    ]
    assert max(lengths) <= 400
    assert all(piece.page_number == 3 for piece in pieces)
    assert all(piece.text in text for piece in pieces)
    assert pieces[1].text.startswith("unaffordability ")
    ids = [real_bge_tokenizer.encode(p.text, add_special_tokens=False).ids for p in pieces]
    assert ids[0][-55:] == ids[1][:55]  # The 50th token is inside this six-token word.
    assert ids[1][-50:] == ids[2][:50]


def test_wordpiece_chunks_stay_bounded_after_independent_encoding():
    tokenizer = Tokenizer(
        models.WordPiece(
            {
                "[UNK]": 0,
                "中": 1,
                "u": 2,
                "##na": 3,
                "##ff": 4,
                "##or": 5,
                "##da": 6,
                "##bility": 7,
                "b": 8,
                "##il": 9,
                "##ity": 10,
            },
            unk_token="[UNK]",
        )
    )
    tokenizer.pre_tokenizer = pre_tokenizers.Whitespace()
    text = "中 " * 345 + "unaffordability " + "中 " * 500
    pieces = split_sections([Section(text)], tokenizer)
    assert all(len(tokenizer.encode(piece.text).ids) <= 400 for piece in pieces)


def test_chunks_preserve_original_unicode_spaces_and_punctuation(real_bge_tokenizer):
    text = "  校园😀  cafe\u0301，\tEnglish 空格。\u200b  "
    pieces = split_sections([Section(text, line_number=7)], real_bge_tokenizer)
    assert [piece.text for piece in pieces] == [text]
    assert pieces[0].line_number == 7


def test_docx_body_order_and_table_locations(tmp_path):
    document = Document()
    document.add_paragraph("表格之前")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "时间"
    table.cell(0, 1).text = "地点"
    table.cell(1, 0).text = "九点"
    table.cell(1, 1).text = "图书馆"
    document.add_paragraph("表格之后")
    path = tmp_path / "ordered.docx"
    document.save(path)
    sections = parse_file(path, "docx")
    assert [section.text for section in sections] == [
        "表格之前",
        "时间 | 地点",
        "九点 | 图书馆",
        "表格之后",
    ]
    assert [(s.paragraph_number, s.line_number) for s in sections] == [
        (1, None),
        (2, 1),
        (2, 2),
        (3, None),
    ]


def test_real_bge_long_unknown_word_terminates_without_losing_text(real_bge_tokenizer):
    word = "a" * 5000
    text = word + " 中" * 450
    pieces = split_sections([Section(text)], real_bge_tokenizer)
    assert len(pieces) == 2
    assert pieces[0].text.startswith(word)
    assert pieces[-1].text.endswith(" 中")
    assert all(
        len(real_bge_tokenizer.encode(p.text, add_special_tokens=False).ids) <= 400 for p in pieces
    )


def test_unrepresentable_single_word_fails_instead_of_looping():
    tokenizer = Tokenizer(
        models.WordPiece(
            {"[UNK]": 0, "a": 1, "##a": 2}, unk_token="[UNK]", max_input_chars_per_word=2000
        )
    )
    tokenizer.pre_tokenizer = pre_tokenizers.Whitespace()
    with pytest.raises(ParseError, match="PARSE_LIMIT"):
        split_sections([Section("a" * 1000)], tokenizer)


def test_chinese_pdf_text_and_page(tmp_path):
    from pypdf.generic import (
        ArrayObject,
        DecodedStreamObject,
        DictionaryObject,
        NameObject,
        NumberObject,
        TextStringObject,
    )

    text = "校园图书馆开放"
    mappings = "\n".join(
        f"<{index:04x}> <{character.encode('utf-16-be').hex()}>"
        for index, character in enumerate(text, 1)
    )
    cmap = DecodedStreamObject()
    cmap.set_data(
        (
            "/CIDInit /ProcSet findresource begin 12 dict begin begincmap\n"
            "/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def\n"
            "/CMapName /Campus def /CMapType 2 def\n"
            "1 begincodespacerange <0000> <FFFF> endcodespacerange\n"
            f"{len(text)} beginbfchar\n{mappings}\nendbfchar\n"
            "endcmap CMapName currentdict /CMap defineresource pop end end"
        ).encode("ascii")
    )
    descendant = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/CIDFontType2"),
            NameObject("/BaseFont"): NameObject("/Campus"),
            NameObject("/CIDSystemInfo"): DictionaryObject(
                {
                    NameObject("/Registry"): TextStringObject("Adobe"),
                    NameObject("/Ordering"): TextStringObject("Identity"),
                    NameObject("/Supplement"): NumberObject(0),
                }
            ),
            NameObject("/DW"): NumberObject(1000),
        }
    )
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type0"),
            NameObject("/BaseFont"): NameObject("/Campus"),
            NameObject("/Encoding"): NameObject("/Identity-H"),
            NameObject("/DescendantFonts"): ArrayObject([descendant]),
            NameObject("/ToUnicode"): writer._add_object(cmap),
        }
    )
    page = writer.add_blank_page(200, 200)
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    codes = "".join(f"{index:04x}" for index in range(1, len(text) + 1))
    stream.set_data(f"BT /F1 12 Tf 10 10 Td <{codes}> Tj ET".encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(stream)
    path = tmp_path / "chinese.pdf"
    writer.write(path)
    sections = parse_file(path, "pdf")
    assert [section.text for section in sections] == [text]
    assert sections[0].page_number == 1


@pytest.mark.parametrize("suffix", ["txt", "md"])
def test_text_formats_keep_indentation_and_trailing_spaces(tmp_path, suffix):
    first = "    print('校园')  "
    third = '\tcode = "a  b"\t'
    path = tmp_path / f"indentation.{suffix}"
    path.write_text(first + "\n\n" + third + "\n", encoding="utf-8")
    sections = parse_file(path, suffix)
    assert [section.text for section in sections] == [first, third]
    assert [section.line_number for section in sections] == [1, 3]


@pytest.mark.parametrize("text", ["\u200b", "\u200c", "\u200d"])
def test_real_bge_zero_token_document_is_no_text(tmp_path, real_bge_tokenizer, text):
    path = tmp_path / "invisible.txt"
    path.write_text(text, encoding="utf-8")
    sections = parse_file(path, "txt")
    assert real_bge_tokenizer.encode(text, add_special_tokens=False).ids == []
    with pytest.raises(ParseError, match="NO_TEXT"):
        split_sections(sections, real_bge_tokenizer)
